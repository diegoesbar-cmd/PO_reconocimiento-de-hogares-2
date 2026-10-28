#!/usr/bin/env python3
"""
Convierte un XLSForm (hojas survey / choices / settings) en cuestionario.js,
el archivo que la app lee para dibujar el cuestionario.

Uso:
    python3 herramientas/convertir_xlsform.py herramientas/cuestionario.xlsx cuestionario.js

Requiere:  pip install openpyxl

Qué hace además de convertir:
  * revisa que cada ${variable} usada en saltos/restricciones/cálculos exista;
  * revisa que cada lista de opciones usada esté definida;
  * avisa de nombres duplicados y de tipos que la app no admite;
  * corrige el error común de poner «multiline» en la columna calculation.
Termina con código 1 si encuentra errores (los avisos no detienen la conversión).
"""
import json
import re
import sys

import openpyxl

# Títulos de pasos para bloques repetidos (si no se indica, se usa la etiqueta del bloque)
TITULOS_REPETIDOS = {
    "integrantes": "Integrantes del hogar",
}
TITULO_PASO_INICIAL = "Datos de la visita"

TIPOS_SOPORTADOS = {
    "start", "end", "text", "integer", "decimal", "date", "time",
    "select_one", "select_multiple", "note", "calculate",
    "begin_group", "end_group", "begin_repeat", "end_repeat",
}

errores, avisos = [], []


def txt(v):
    if v is None:
        return ""
    return str(v).strip()


def si(v):
    return txt(v).lower() in ("yes", "true", "1", "si", "sí")


def leer_hoja(wb, nombre):
    if nombre not in wb.sheetnames:
        errores.append(f"Falta la hoja «{nombre}».")
        return []
    ws = wb[nombre]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    cab = [txt(c) for c in filas[0]]
    out = []
    for n, f in enumerate(filas[1:], start=2):
        d = {cab[i]: f[i] for i in range(len(cab)) if cab[i]}
        d["_fila"] = n
        if any(txt(v) for k, v in d.items() if k != "_fila"):
            out.append(d)
    return out


def main(origen, destino):
    wb = openpyxl.load_workbook(origen, data_only=True)
    survey = leer_hoja(wb, "survey")
    choices = leer_hoja(wb, "choices")
    settings = leer_hoja(wb, "settings")

    # ---------------- listas de opciones
    listas = {}
    for r in choices:
        ln, nm = txt(r.get("list_name")), txt(r.get("name"))
        if not ln:
            continue
        listas.setdefault(ln, []).append([nm, txt(r.get("label"))])

    # ---------------- árbol de preguntas
    raiz = []
    pila = [("raiz", raiz)]
    nombres = {}
    meta_auto = {"start": None, "end": None}

    for r in survey:
        fila = r["_fila"]
        tipo_raw = txt(r.get("type"))
        if not tipo_raw:
            continue
        partes = tipo_raw.split()
        tipo = partes[0]
        nombre = txt(r.get("name"))

        if tipo not in TIPOS_SOPORTADOS:
            errores.append(f"Fila {fila}: tipo «{tipo_raw}» no soportado por la app.")
            continue

        if tipo in ("end_group", "end_repeat"):
            if len(pila) == 1:
                errores.append(f"Fila {fila}: {tipo} sin su begin correspondiente.")
            else:
                pila.pop()
            continue

        if tipo in ("start", "end"):
            meta_auto[tipo] = nombre
            continue

        if not nombre:
            errores.append(f"Fila {fila}: falta la columna name.")
            continue
        if nombre in nombres:
            errores.append(f"Fila {fila}: el nombre «{nombre}» está repetido (también en la fila {nombres[nombre]}).")
        nombres[nombre] = fila

        apariencia = txt(r.get("appearance"))
        calculo = txt(r.get("calculation"))
        if tipo != "calculate" and calculo in ("multiline", "minimal", "field-list"):
            avisos.append(
                f"Fila {fila} ({nombre}): «{calculo}» está en la columna calculation; "
                f"se trata como appearance. Conviene corregirlo en el Excel."
            )
            if not apariencia:
                apariencia = calculo
            calculo = ""

        item = {"name": nombre, "type": tipo, "label": txt(r.get("label"))}
        if txt(r.get("hint")):
            item["hint"] = txt(r.get("hint"))
        if si(r.get("required")):
            item["required"] = True
        if txt(r.get("relevant")):
            item["relevant"] = txt(r.get("relevant"))
        if txt(r.get("constraint")):
            item["constraint"] = txt(r.get("constraint"))
        if txt(r.get("constraint_message")):
            item["constraintMsg"] = txt(r.get("constraint_message"))
        if apariencia == "multiline":
            item["multiline"] = True
        item["_fila"] = fila

        if tipo in ("select_one", "select_multiple"):
            if len(partes) < 2:
                errores.append(f"Fila {fila} ({nombre}): falta el nombre de la lista de opciones.")
                continue
            item["list"] = partes[1]
            if partes[1] not in listas:
                errores.append(f"Fila {fila} ({nombre}): la lista «{partes[1]}» no existe en la hoja choices.")
        elif tipo == "calculate":
            if not calculo:
                errores.append(f"Fila {fila} ({nombre}): calculate sin expresión.")
            item["calc"] = calculo
        elif tipo in ("begin_group", "begin_repeat"):
            item["type"] = "group" if tipo == "begin_group" else "repeat"
            item["items"] = []

        pila[-1][1].append(item)
        if tipo in ("begin_group", "begin_repeat"):
            pila.append((item["type"], item["items"]))

    if len(pila) != 1:
        errores.append("Hay un begin_group/begin_repeat sin cerrar.")

    # ---------------- referencias ${...}
    def revisar_refs(items):
        for it in items:
            for campo in ("relevant", "constraint", "calc"):
                for ref in re.findall(r"\$\{([^}]+)\}", it.get(campo, "")):
                    if ref not in nombres:
                        errores.append(f"Fila {it['_fila']} ({it['name']}): {campo} usa ${{{ref}}}, que no existe.")
            if "items" in it:
                revisar_refs(it["items"])

    revisar_refs(raiz)

    # ---------------- pasos: cada grupo de primer nivel es un paso;
    # lo que queda suelto entre grupos (incluido un bloque repetido) forma otro paso.
    pasos = []
    suelto = []

    def vaciar_suelto():
        nonlocal suelto
        if not suelto:
            return
        rep = next((x for x in suelto if x["type"] == "repeat"), None)
        if rep:
            titulo = TITULOS_REPETIDOS.get(rep["name"], rep["label"])
        elif not pasos:
            titulo = TITULO_PASO_INICIAL
        else:
            titulo = "Continuación"
        pasos.append({"id": f"paso{len(pasos) + 1}", "titulo": titulo, "items": suelto})
        suelto = []

    for it in raiz:
        if it["type"] == "group":
            vaciar_suelto()
            paso = {"id": it["name"], "titulo": it["label"], "items": it["items"]}
            if it.get("relevant"):
                paso["relevant"] = it["relevant"]
            if it.get("hint"):
                paso["hint"] = it["hint"]
            pasos.append(paso)
        else:
            suelto.append(it)
            if it["type"] == "repeat":
                vaciar_suelto()
    vaciar_suelto()

    def limpiar(items):
        for it in items:
            it.pop("_fila", None)
            if "items" in it:
                limpiar(it["items"])

    for p in pasos:
        limpiar(p["items"])

    cfg = settings[0] if settings else {}
    salida = {
        "meta": {
            "titulo": txt(cfg.get("form_title")) or "Cuestionario",
            "formId": txt(cfg.get("form_id")) or "cuestionario",
            "version": txt(cfg.get("version")) or "1",
            "inicio": meta_auto["start"],
            "fin": meta_auto["end"],
        },
        "listas": listas,
        "pasos": pasos,
    }

    cuerpo = json.dumps(salida, ensure_ascii=False, indent=1)
    # pares [valor, etiqueta] en una sola línea, para que el archivo sea legible
    cuerpo = re.sub(
        r'\[\s+("(?:[^"\\]|\\.)*"),\s+("(?:[^"\\]|\\.)*")\s+\]', r"[\1, \2]", cuerpo
    )
    js = (
        "/* GENERADO por herramientas/convertir_xlsform.py a partir del XLSForm. No editar a mano. */\n"
        "window.CUESTIONARIO = " + cuerpo + ";\n"
    )
    for a in avisos:
        print("AVISO:", a)
    for e in errores:
        print("ERROR:", e)
    if errores:
        sys.exit(1)
    with open(destino, "w", encoding="utf-8") as f:
        f.write(js)

    def contar(items):
        n = 0
        for it in items:
            if it["type"] not in ("group", "repeat"):
                n += 1
            if "items" in it:
                n += contar(it["items"])
        return n

    total = sum(contar(p["items"]) for p in pasos)
    print(f"OK: {len(pasos)} pasos, {total} preguntas/cálculos/notas, {len(listas)} listas -> {destino}")
    for p in pasos:
        print(f"  - {p['titulo']} ({contar(p['items'])})")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
